"""Document Bidirectional Conversion Test Suite (PDF <-> Word, Excel, PowerPoint, Text, HTML)"""
import pytest
import io
from docx import Document
from openpyxl import Workbook
from pptx import Presentation


@pytest.fixture
def sample_xlsx_bytes():
    """Generates an in-memory XLSX workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "PaperKit Sheet"
    ws.append(["ID", "Name", "Score"])
    ws.append([1, "Item A", 95.5])
    ws.append([2, "Item B", 88.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_pptx_bytes():
    """Generates an in-memory PPTX presentation."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "PaperKit Slide"
    subtitle.text = "Conversion testing slide"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_convert_pdf_to_word(client, seeded_file, auth_headers):
    """POST /tools/convert converts PDF to Word DOCX."""
    file_id = str(seeded_file["_id"])
    resp = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "pdf", "to_format": "word"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert "download_url" in data
    assert data["download_url"].endswith(".docx")


@pytest.mark.asyncio
async def test_convert_pdf_to_excel(client, seeded_file, auth_headers):
    """POST /tools/convert converts PDF to Excel XLSX."""
    file_id = str(seeded_file["_id"])
    resp = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "pdf", "to_format": "excel"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert "download_url" in data
    assert data["download_url"].endswith(".xlsx")


@pytest.mark.asyncio
async def test_convert_pdf_to_ppt(client, seeded_file, auth_headers):
    """POST /tools/convert converts PDF to PPTX."""
    file_id = str(seeded_file["_id"])
    resp = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "pdf", "to_format": "ppt"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert "download_url" in data
    assert data["download_url"].endswith(".pptx")


@pytest.mark.asyncio
async def test_convert_word_to_pdf(client, sample_docx_bytes, auth_headers):
    """POST /tools/convert converts DOCX to PDF."""
    files = {"file": ("document.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    up_resp = await client.post("/files/upload", files=files, headers=auth_headers)
    file_id = up_resp.json()["_id"]

    resp = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "word", "to_format": "pdf"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.json()["download_url"].endswith(".pdf")


@pytest.mark.asyncio
async def test_convert_excel_to_pdf(client, sample_xlsx_bytes, auth_headers):
    """POST /tools/convert converts XLSX to PDF."""
    files = {"file": ("sheet.xlsx", sample_xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    up_resp = await client.post("/files/upload", files=files, headers=auth_headers)
    file_id = up_resp.json()["_id"]

    resp = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "excel", "to_format": "pdf"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.json()["download_url"].endswith(".pdf")


@pytest.mark.asyncio
async def test_convert_ppt_to_pdf(client, sample_pptx_bytes, auth_headers):
    """POST /tools/convert converts PPTX to PDF."""
    files = {"file": ("presentation.pptx", sample_pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    up_resp = await client.post("/files/upload", files=files, headers=auth_headers)
    file_id = up_resp.json()["_id"]

    resp = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "ppt", "to_format": "pdf"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.json()["download_url"].endswith(".pdf")


@pytest.mark.asyncio
async def test_convert_html_to_word_endpoint(client, auth_headers):
    """POST /tools/html-to-word converts HTML string directly to Word DOCX."""
    resp = await client.post(
        "/tools/html-to-word",
        json={"html_content": "<p>PaperKit Generated Paragraph</p>", "filename": "output.docx"},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.json()["download_url"].endswith(".docx")


@pytest.fixture
def sample_pdf_with_table_and_image():
    """Generates an in-memory PDF containing a structured table, borderless table, and image."""
    from PIL import Image as PILImage
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, Image
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    img = PILImage.new("RGB", (100, 60), color=(59, 130, 246))
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_buf.seek(0)

    pdf_buf = io.BytesIO()
    doc = SimpleDocTemplate(pdf_buf, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Document with Tables and Graphics", styles["Heading1"]),
        Spacer(1, 10),
        Paragraph("Summary section before table.", styles["Normal"]),
        Spacer(1, 10),
        Table([
            ["Item", "Quantity", "Unit Price", "Total"],
            ["Server License", "2", "$500", "$1000"],
            ["Support Tier", "1", "$250", "$250"],
        ], style=[
            ("GRID", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ]),
        Spacer(1, 12),
        Image(img_buf, width=100, height=60),
        Spacer(1, 12),
        Table([
            ["Borderless Header A", "Borderless Header B"],
            ["Data Value 1", "Data Value 2"],
        ]),
        Spacer(1, 10),
        Paragraph("End of document.", styles["Normal"]),
    ]
    doc.build(story)
    return pdf_buf.getvalue()


@pytest.mark.asyncio
async def test_pdf_with_tables_and_images_conversions(client, sample_pdf_with_table_and_image, auth_headers):
    """Verify table and image detection across Word, Excel, PPT, and HTML conversion flows."""
    files = {"file": ("report_with_media.pdf", sample_pdf_with_table_and_image, "application/pdf")}
    up_resp = await client.post("/files/upload", files=files, headers=auth_headers)
    assert up_resp.status_code == 200
    file_id = up_resp.json()["_id"]

    # 1. Convert to Word
    resp_word = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "pdf", "to_format": "word"},
        headers=auth_headers,
    )
    assert resp_word.status_code == 200
    assert resp_word.json()["download_url"].endswith(".docx")
    assert resp_word.json()["size"] > 0

    # 2. Convert to Excel
    resp_excel = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "pdf", "to_format": "excel"},
        headers=auth_headers,
    )
    assert resp_excel.status_code == 200
    assert resp_excel.json()["download_url"].endswith(".xlsx")
    assert resp_excel.json()["size"] > 0

    # 3. Convert to PPT
    resp_ppt = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "pdf", "to_format": "ppt"},
        headers=auth_headers,
    )
    assert resp_ppt.status_code == 200
    assert resp_ppt.json()["download_url"].endswith(".pptx")
    assert resp_ppt.json()["size"] > 0

    # 4. Convert to HTML
    resp_html = await client.post(
        "/tools/convert",
        json={"file_id": file_id, "from_format": "pdf", "to_format": "html"},
        headers=auth_headers,
    )
    assert resp_html.status_code == 200
    assert resp_html.json()["html_content"]
    assert "<table" in resp_html.json()["html_content"]
    assert "<img" in resp_html.json()["html_content"]

