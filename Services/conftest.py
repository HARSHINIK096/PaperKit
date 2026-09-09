"""PaperKit Test Suite — Central Fixtures and Configurations
Provides isolated, fast test execution with in-memory database and async FastAPI TestClient.
"""
import os
import sys
import pytest
import io
import fitz  # PyMuPDF
from datetime import timezone, datetime
from bson import ObjectId
from docx import Document
from PIL import Image

# Ensure the Services directory is on sys.path
SERVICES_DIR = os.path.dirname(__file__)
if SERVICES_DIR not in sys.path:
    sys.path.insert(0, SERVICES_DIR)

# Configure environment variables for test execution
os.environ["MONGODB_URL"] = "mock://"
os.environ["DATABASE_NAME"] = "paperkit_test"
os.environ["SECRET_KEY"] = "test-paperkit-jwt-secret-key-32-chars-long"
os.environ["ALGORITHM"] = "HS256"
os.environ["ACCESS_TOKEN_EXPIRE_MINUTES"] = "120"
os.environ["GEMINI_API_KEY"] = "mock-gemini-key"
os.environ["GROQ_API_KEY"] = "mock-groq-key"
os.environ["FRONTEND_URL"] = "https://paperkit-web.onrender.com"

from config import get_settings
get_settings.cache_clear()

from httpx import AsyncClient, ASGITransport
from middleware.auth import create_access_token, hash_password
from database import get_db, MockDatabase, MockDatabaseClient
import database as db_module
from main import app as fastapi_app


@pytest.fixture(autouse=True)
def reset_db():
    """Reset the mock database to a clean in-memory state before each test."""
    fresh_db = MockDatabase.__new__(MockDatabase)
    fresh_db.file_path = None
    fresh_db._data = {}
    fresh_db.save_db = lambda: None

    fresh_client = MockDatabaseClient.__new__(MockDatabaseClient)
    fresh_client.database = fresh_db
    db_module._client = fresh_client
    yield
    db_module._client = None


@pytest.fixture
def db():
    """Returns the mock database instance."""
    return get_db()


@pytest.fixture
def test_user_doc():
    """Pre-built registered user document."""
    user_id = ObjectId()
    return {
        "_id": user_id,
        "name": "PaperKit Tester",
        "email": "tester@paperkit.dev",
        "hashed_password": hash_password("TestPassword123!"),
        "avatar_url": None,
        "oauth_provider": None,
        "oauth_id": None,
        "created_at": datetime.now(timezone.utc),
        "preferences": {
            "dark_mode": True,
            "default_view": "grid",
            "language": "en",
        },
    }


@pytest.fixture
def auth_token(test_user_doc):
    """Generates a valid JWT bearer token for the test user."""
    return create_access_token({"sub": str(test_user_doc["_id"])})


@pytest.fixture
def auth_headers(seeded_user, auth_token):
    """Returns Authorization header with Bearer token for seeded test user."""
    return {"Authorization": f"Bearer {auth_token}"}


@pytest.fixture
async def seeded_user(db, test_user_doc):
    """Inserts the test user into DB and returns user doc."""
    await db.users.insert_one(test_user_doc)
    return test_user_doc


@pytest.fixture
def sample_pdf_bytes():
    """Generates a real valid PDF in memory with selectable text."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "PaperKit Universal Test Document", fontsize=16)
    page.insert_text((72, 140), "This document tests PDF manipulation, encryption, OCR, and AI operations.", fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def multi_page_pdf_bytes():
    """Generates a 3-page PDF for split, organize, and page extraction tests."""
    doc = fitz.open()
    for i in range(1, 4):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 100), f"PaperKit Test Page {i}", fontsize=18)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_image_bytes():
    """Generates a sample PNG image in memory."""
    img = Image.new("RGB", (200, 200), color=(79, 70, 229))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def sample_docx_bytes():
    """Generates a valid DOCX file in memory."""
    doc = Document()
    doc.add_heading("PaperKit Test DOCX", level=1)
    doc.add_paragraph("This is paragraph content for bidirectional document conversion tests.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
async def seeded_file(db, seeded_user, sample_pdf_bytes):
    """Creates a physical file in storage and adds a corresponding DB record."""
    from services.storage import LOCAL_STORAGE_DIR
    os.makedirs(LOCAL_STORAGE_DIR, exist_ok=True)
    file_id = ObjectId()
    filename = f"test_{file_id}.pdf"
    file_path = os.path.join(LOCAL_STORAGE_DIR, filename)
    with open(file_path, "wb") as f:
        f.write(sample_pdf_bytes)

    file_doc = {
        "_id": file_id,
        "user_id": str(seeded_user["_id"]),
        "original_filename": "test_document.pdf",
        "content_type": "application/pdf",
        "size": len(sample_pdf_bytes),
        "page_count": 1,
        "storage_url": f"/storage/{filename}",
        "is_deleted": False,
        "created_at": datetime.now(timezone.utc),
        "updated_at": datetime.now(timezone.utc),
    }
    await db.files.insert_one(file_doc)
    yield file_doc
    if os.path.exists(file_path):
        try:
            os.remove(file_path)
        except OSError:
            pass


@pytest.fixture
def app():
    """Returns the FastAPI application instance."""
    return fastapi_app


@pytest.fixture
async def client(app):
    """Async HTTP test client."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac
